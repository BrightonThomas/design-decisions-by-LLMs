from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from . import SCHEMA_VERSION
from .core import write_jsonl


DEFAULT_PILOT_TRIALS = {
    "c68f733c-d7d7-52fa-b063-e3b5cbd29eef",  # Claude Fable / Claude Code
    "30d171a7-5d7a-50b1-8d7d-31e8e8b6a734",  # Muse Spark / mini-swe-agent
}


def records_from_sheet(workbook_path: Path, sheet_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook[sheet_name]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value) for value in next(rows)]
    return [dict(zip(headers, row)) for row in rows if any(value is not None for value in row)]


def task_description_from_source(source_workbook: Path) -> str:
    workbook = load_workbook(source_workbook, read_only=True, data_only=True)
    return str(workbook.active.cell(row=2, column=1).value or "")


def expand_steps(value: Any) -> list[int]:
    result: set[int] = set()
    for token in re.split(r"[,;\s]+", str(value or "").strip()):
        if not token:
            continue
        if "-" in token:
            start, end = token.split("-", 1)
            result.update(range(int(start), int(end) + 1))
        else:
            result.add(int(token))
    return sorted(result)


def short_exact_excerpt(value: Any, maximum: int = 600) -> str | None:
    if not isinstance(value, str) or not value.strip():
        return None
    return value[:maximum]


def build_assets(gold_workbook: Path, source_workbook: Path) -> tuple[list[dict], list[dict]]:
    step_rows = records_from_sheet(gold_workbook, "Trajectory Steps")
    event_rows = records_from_sheet(gold_workbook, "Decision Events")
    task_description = task_description_from_source(source_workbook)

    steps_by_trial: dict[str, list[dict[str, Any]]] = defaultdict(list)
    metadata: dict[str, dict[str, Any]] = {}
    for row in step_rows:
        trial_id = str(row["Trial ID"])
        metadata[trial_id] = {
            "task_id": row["Task ID"], "trial_id": trial_id,
            "model": row["Model"], "harness": row["Harness"],
            "resolved": bool(row["Resolved"]),
        }
        steps_by_trial[trial_id].append({
            "index": int(row["Step"]),
            "thought": row["Thought"],
            "action": row["Action"],
            "observation": row["Observation"],
        })

    trajectories = []
    for trial_id, info in metadata.items():
        trajectories.append({
            **info,
            "task_description": task_description,
            "steps": sorted(steps_by_trial[trial_id], key=lambda item: item["index"]),
        })

    events_by_trial: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in event_rows:
        trial_id = str(row["Trial ID"])
        source_steps = expand_steps(row["Source steps"])
        source_lookup = {step["index"]: step for step in steps_by_trial[trial_id]}
        evidence = []
        for field in ("thought", "action", "observation"):
            for step_id in source_steps:
                excerpt = short_exact_excerpt(source_lookup[step_id].get(field))
                if excerpt:
                    evidence.append({"step": step_id, "field": field, "excerpt": excerpt})
                    break
        events_by_trial[trial_id].append({
            "event_id": str(row["Event ID"]),
            "source_steps": source_steps,
            "primary_category": row["Primary category"],
            "secondary_category": row["Secondary category"] or None,
            "decision_summary": row["Decision summary"],
            "selected_option": row["Selected option"],
            "alternative_visible": row["Alternative visible"] or None,
            "primary_reason": row["Primary reason"],
            "secondary_reasons": [
                part.strip() for part in str(row["Secondary reasons"] or "").split(";") if part.strip()
            ],
            "reason_summary": "" if row["Reason visibility"] == "UNAVAILABLE" else (row["Reason summary"] or ""),
            "decision_visibility": row["Decision visibility"],
            "reason_visibility": row["Reason visibility"],
            "confidence": row["Confidence"],
            "evidence": evidence,
        })

    gold = [{
        "schema_version": SCHEMA_VERSION,
        "task_id": metadata[trial_id]["task_id"],
        "trial_id": trial_id,
        "events": events,
    } for trial_id, events in events_by_trial.items()]
    return trajectories, gold


def replay_fixture(gold_record: dict[str, Any]) -> dict[str, Any]:
    response = {
        **gold_record,
        "events": [{**event, "event_id": f"P{i:02d}"} for i, event in enumerate(gold_record["events"], 1)],
    }
    return {"trial_id": gold_record["trial_id"], "response": response}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export Week 1 gold data for the Week 2 tool")
    parser.add_argument("--gold-workbook", type=Path, required=True)
    parser.add_argument("--source-workbook", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    trajectories, gold = build_assets(args.gold_workbook, args.source_workbook)
    pilot_trajectories = [item for item in trajectories if item["trial_id"] in DEFAULT_PILOT_TRIALS]
    pilot_gold = [item for item in gold if item["trial_id"] in DEFAULT_PILOT_TRIALS]
    args.output_dir.mkdir(parents=True, exist_ok=True)
    write_jsonl(trajectories, args.output_dir / "trajectories_all6.jsonl")
    write_jsonl(gold, args.output_dir / "gold_events_all6.jsonl")
    write_jsonl((replay_fixture(item) for item in gold), args.output_dir / "gold_replay_fixture_all6.jsonl")
    write_jsonl(pilot_trajectories, args.output_dir / "pilot_trajectories_two.jsonl")
    write_jsonl(pilot_gold, args.output_dir / "pilot_gold_events_two.jsonl")
    write_jsonl((replay_fixture(item) for item in pilot_gold), args.output_dir / "gold_replay_fixture_two.jsonl")
    print({
        "trajectories": len(trajectories),
        "gold_events": sum(len(item["events"]) for item in gold),
        "pilot_trajectories": len(pilot_trajectories),
        "pilot_events": sum(len(item["events"]) for item in pilot_gold),
    })


if __name__ == "__main__":
    main()
